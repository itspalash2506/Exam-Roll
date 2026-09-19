"""Excel formula-injection guard (P0-5, DECISIONS.md 2026-09-19).

Before workbook_builder.py existed, every value written to a cell in
excel_generator.py was unescaped: a roll number, subject code/name, or
AI-derived exam_name/course/semester starting with `=`, `+`, `-` or `@`
became a LIVE FORMULA the moment the delivered workbook was opened. These
tests build a payload deliberately shaped like an attack and assert every
untrusted sink is neutralized, while the generator's OWN formulas (COUNTA
totals, SUM total) remain real, working formulas.
"""

import io

import openpyxl

from app.schemas.schemas import ExtractedDataSchema, StudentRecord, StyleConfig, SubjectEntry
from app.services.generators.excel_generator import generate_excel
from app.services.generators.workbook_builder import safe_value


def test_safe_value_quote_prefixes_formula_triggers():
    for trigger in ("=", "+", "-", "@", "\t", "\r"):
        payload = f"{trigger}HYPERLINK(\"http://evil\")"
        assert safe_value(payload) == "'" + payload


def test_safe_value_passes_through_non_triggering_and_non_string():
    assert safe_value("10001") == "10001"
    assert safe_value(None) is None
    assert safe_value(42) == 42
    assert safe_value("") == ""


def _malicious_payload() -> ExtractedDataSchema:
    subjects = [
        # code AND name both start with formula-trigger characters
        SubjectEntry(code="=HYPERLINK(\"http://evil\")", name="+cmd|'/c calc'!A1"),
    ]
    students = [
        StudentRecord(roll_number="=1+1", subjects=[subjects[0].code]),
        StudentRecord(roll_number="10002", subjects=[subjects[0].code]),
    ]
    return ExtractedDataSchema(
        students=students,
        subjects=subjects,
        source_file="attack.pdf",
        total_students=len(students),
        document_type="attestation_sheet",
        course="@SUM(A1:A100)",
        semester="-2+3",
        exam_name="=cmd|'/c calc'!A1",
        ai_confidence=0.5,
    )


def _load_workbook(xlsx_bytes: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(xlsx_bytes))


def test_every_untrusted_sink_is_neutralized():
    data = _malicious_payload()
    xlsx_bytes = generate_excel(data, StyleConfig(), "attack-output")
    wb = _load_workbook(xlsx_bytes)

    ws1 = wb["Subject-wise Roll Number List"]
    ws2 = wb["Summary"]

    # Sheet 1 title (row 1) embeds exam_name/course/semester, but is always
    # prefixed with the literal "SUBJECT-WISE ROLL NUMBER LIST" — so the
    # embedded malicious text is never at the START of the cell value, and
    # Excel only treats a LEADING =/+/-/@ as a formula trigger. Correctly
    # not quote-prefixed by safe_value(); what matters is it never became a
    # formula cell.
    title_cell = ws1.cell(row=1, column=1)
    assert title_cell.data_type != "f", "title cell became a live formula"

    # Sheet 1 subject header (row 2, col 2): "code\nname", both malicious.
    header_cell = ws1.cell(row=2, column=2)
    assert header_cell.data_type != "f"
    assert header_cell.value.startswith("'")

    # Both roll values in column B (rows 3-4: two students share the one
    # malicious subject; export-time sort order between "=1+1" and "10002"
    # is not the point of this test, so check both data rows rather than
    # assume which row holds which value).
    roll_cells = [ws1.cell(row=r, column=2) for r in (3, 4)]
    roll_values = {c.value for c in roll_cells}
    assert all(c.data_type != "f" for c in roll_cells), "a roll number became a live formula"
    assert "'=1+1" in roll_values, "malicious roll number not quote-prefixed"
    assert "10002" in roll_values, "benign roll number was altered"

    # Sheet 1 count row's COUNTA formula must remain a REAL, working formula.
    # data_rows = max(rolls per subject) = 2 here, so the count row is row 5.
    count_cell = ws1.cell(row=5, column=2)
    assert count_cell.data_type == "f", "generator's own COUNTA formula was sanitized away"
    assert count_cell.value.startswith("=COUNTA(")

    # Sheet 2 metadata block: course/semester/exam_name values (col 2).
    for row in range(2, 5):  # Document Type(2), Course(3), Semester(4)
        pass
    course_cell = ws2.cell(row=3, column=2)
    semester_cell = ws2.cell(row=4, column=2)
    exam_name_cell = ws2.cell(row=5, column=2)
    for cell, label in [(course_cell, "course"), (semester_cell, "semester"), (exam_name_cell, "exam_name")]:
        assert cell.data_type != "f", f"{label} cell became a live formula"
        assert cell.value.startswith("'"), f"{label} cell not quote-prefixed"

    # Sheet 2 summary table row: subject code (col 2) and name (col 3).
    code_cell = ws2.cell(row=10, column=2)
    name_cell = ws2.cell(row=10, column=3)
    assert code_cell.data_type != "f"
    assert code_cell.value.startswith("'")
    assert name_cell.data_type != "f"
    assert name_cell.value.startswith("'")

    # Sheet 2 SUM total (1 subject -> total row is row 11, col 4) must
    # remain a REAL, working formula.
    sum_cell = ws2.cell(row=11, column=4)
    assert sum_cell.data_type == "f", "generator's own SUM formula was sanitized away"
    assert sum_cell.value.startswith("=SUM(")


def test_benign_payload_is_untouched():
    """A normal, non-malicious payload should read back exactly as before —
    the guard must be invisible in the common case."""
    subjects = [SubjectEntry(code="MBAN301", name="Business Mathematics")]
    students = [StudentRecord(roll_number="10001", subjects=["MBAN301"])]
    data = ExtractedDataSchema(
        students=students,
        subjects=subjects,
        source_file="test.pdf",
        total_students=1,
        document_type="attestation_sheet",
        course="B.Com",
        semester="3rd Semester",
        exam_name="Nov 2024",
        ai_confidence=0.9,
    )
    xlsx_bytes = generate_excel(data, StyleConfig(), "benign-output")
    wb = _load_workbook(xlsx_bytes)
    ws1 = wb["Subject-wise Roll Number List"]

    roll_cell = ws1.cell(row=3, column=2)
    assert roll_cell.value == "10001"  # no stray quote prefix
    header_cell = ws1.cell(row=2, column=2)
    assert header_cell.value == "MBAN301\nBusiness Mathematics"
