"""Golden-file regression test for the Subject-wise Excel generator (P11).

tests/golden/subject_wise.xlsx is generated once from _fixture_extracted_data()
below (see scripts/generate_golden_subject_wise.py) and committed. This test
regenerates from the SAME fixture on every run and compares cell values,
number formats, merged ranges, and print setup against the committed file —
a change here means the generator's actual output changed, which should
never happen silently.
"""

import io
from pathlib import Path

import openpyxl

from app.schemas.schemas import ExtractedDataSchema, StudentRecord, StyleConfig, SubjectEntry
from app.services.generators.excel_generator import generate_excel

_GOLDEN_PATH = Path(__file__).parent / "golden" / "subject_wise.xlsx"


def _fixture_extracted_data() -> ExtractedDataSchema:
    return ExtractedDataSchema(
        students=[
            StudentRecord(roll_number="00123", subjects=["MBAN301", "MBAN302"]),
            StudentRecord(roll_number="10045", subjects=["MBAN301"]),
            StudentRecord(roll_number="10046", subjects=["MBAN302"]),
        ],
        subjects=[
            SubjectEntry(code="MBAN301", name="Business Mathematics"),
            SubjectEntry(code="MBAN302", name="Accountancy"),
        ],
        source_file="golden_fixture.pdf",
        total_students=3,
        document_type="attestation_sheet",
        course="Master of Business Administration",
        semester="3rd Semester",
        exam_name="End Term Examination 2026",
        ai_confidence=0.95,
    )


def _generate_fixture_bytes() -> bytes:
    return generate_excel(_fixture_extracted_data(), StyleConfig(), "Golden Fixture")


def test_golden_fixture_exists():
    assert _GOLDEN_PATH.exists(), (
        "tests/golden/subject_wise.xlsx is missing — regenerate with "
        "scripts/generate_golden_subject_wise.py and commit it."
    )


def test_subject_wise_matches_golden_file():
    generated = openpyxl.load_workbook(io.BytesIO(_generate_fixture_bytes()))
    golden = openpyxl.load_workbook(_GOLDEN_PATH)

    # Summary!B7 is "Generated On" — a real timestamp, genuinely different on
    # every run by design, not a regression signal.
    _SKIP_VALUE_CELLS = {("Summary", 7, 2)}

    for sheet_name in ("Subject-wise Roll Number List", "Summary"):
        gen_ws = generated[sheet_name]
        gold_ws = golden[sheet_name]

        # Cell values across the full used range.
        for row in range(1, max(gen_ws.max_row, gold_ws.max_row) + 1):
            for col in range(1, max(gen_ws.max_column, gold_ws.max_column) + 1):
                if (sheet_name, row, col) in _SKIP_VALUE_CELLS:
                    continue
                gen_cell = gen_ws.cell(row=row, column=col)
                gold_cell = gold_ws.cell(row=row, column=col)
                assert gen_cell.value == gold_cell.value, (
                    f"{sheet_name}!{gen_cell.coordinate}: "
                    f"{gen_cell.value!r} != golden {gold_cell.value!r}"
                )
                assert gen_cell.number_format == gold_cell.number_format, (
                    f"{sheet_name}!{gen_cell.coordinate} number_format: "
                    f"{gen_cell.number_format!r} != golden {gold_cell.number_format!r}"
                )

        # Merged cell ranges.
        assert sorted(str(r) for r in gen_ws.merged_cells.ranges) == sorted(
            str(r) for r in gold_ws.merged_cells.ranges
        ), f"{sheet_name}: merged ranges differ"

        # Print setup.
        assert gen_ws.page_setup.orientation == gold_ws.page_setup.orientation
        assert gen_ws.page_setup.paperSize == gold_ws.page_setup.paperSize
        assert gen_ws.page_setup.fitToWidth == gold_ws.page_setup.fitToWidth
        assert gen_ws.print_area == gold_ws.print_area
        assert gen_ws.print_title_rows == gold_ws.print_title_rows


def test_roll_numbers_are_ascending_and_text_formatted():
    """Sanity check on the fixture itself, independent of the golden file —
    catches the fixture and the golden file silently drifting together."""
    wb = openpyxl.load_workbook(io.BytesIO(_generate_fixture_bytes()))
    ws = wb["Subject-wise Roll Number List"]

    # Column B is MBAN301 (00123, 10045); column C is MBAN302 (00123, 10046).
    assert ws.cell(row=3, column=2).value == "00123"
    assert ws.cell(row=4, column=2).value == "10045"
    assert ws.cell(row=3, column=2).number_format == "@"
