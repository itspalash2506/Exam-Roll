"""Excel export sources from the relational model when it's populated
(P11, DECISIONS.md 2026-09-20) — not the legacy students_json/subjects_json
blob. Reuses the exam+college upload helpers from test_exam_model.py."""

import io

import openpyxl
from sqlalchemy import select

from app.models.db_models import Job, OutputFile, SubjectOffering
from app.routers.export import _load_extracted_relational
from tests.test_exam_model import _create_college, _create_exam, _org_id, _upload


async def _get_job(test_session_factory, job_id):
    async with test_session_factory() as session:
        result = await session.execute(select(Job).where(Job.id == job_id))
        return result.scalar_one()


async def test_load_extracted_relational_returns_none_without_exam_id(
    client, test_session_factory, make_pdf_pages
):
    """A job that predates the exam/college picker has no relational rows —
    the loader must return None so the caller falls back to the blob, not
    raise or produce an empty export."""
    pages = ["MBAN301 - Business Mathematics\nRoll No: 30001  Regular\n"]
    job_id = await _upload(client, None, None, pages, make_pdf_pages, filename="no_exam2.pdf")
    job = await _get_job(test_session_factory, job_id)

    async with test_session_factory() as session:
        result = await _load_extracted_relational(job, session)

    assert result is None


async def test_load_extracted_relational_sources_real_rows(
    client, test_session_factory, make_pdf_pages
):
    exam_id = await _create_exam(client, title="Relational Export Exam")
    college_id = await _create_college(client, name="Relational Export College")
    pages = [
        "MBAN301 - Business Mathematics  MBAN302 - Accountancy\n"
        "Roll No: 40001  Priya S  Regular\n"
        "Roll No: 40002  Arjun K  Regular\n"
    ]
    job_id = await _upload(client, exam_id, college_id, pages, make_pdf_pages, filename="rel.pdf")
    job = await _get_job(test_session_factory, job_id)
    assert job.exam_id == exam_id

    async with test_session_factory() as session:
        extracted = await _load_extracted_relational(job, session)

    assert extracted is not None
    assert {s.roll_number for s in extracted.students} == {"40001", "40002"}
    assert {s.code for s in extracted.subjects} == {"MBAN301", "MBAN302"}
    # Column order must be numeric-suffix ascending, same rule the blob path uses.
    assert [s.code for s in extracted.subjects] == ["MBAN301", "MBAN302"]


async def test_export_reflects_a_later_relational_correction(
    client, test_session_factory, make_pdf_pages
):
    """Proves export genuinely reads the relational rows at export time, not
    a snapshot from upload time — edit a SubjectOffering's name directly in
    the DB after upload and confirm the exported workbook shows the edit."""
    exam_id = await _create_exam(client, title="Correction Exam")
    college_id = await _create_college(client, name="Correction College")
    pages = ["MBAN301 - Biz Math (typo)\nRoll No: 50001  Regular\n"]
    job_id = await _upload(client, exam_id, college_id, pages, make_pdf_pages, filename="corr.pdf")

    async with test_session_factory() as session:
        result = await session.execute(
            select(SubjectOffering).where(SubjectOffering.exam_id == exam_id)
        )
        offering = result.scalar_one()
        offering.subject_name = "Business Mathematics (corrected)"
        await session.commit()

    res = await client.post("/api/v1/export", json={"job_id": job_id, "filename": "corrected"})
    assert res.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb["Subject-wise Roll Number List"]
    header = ws.cell(row=2, column=2).value
    assert "Business Mathematics (corrected)" in header


async def test_roll_number_cells_are_formatted_as_text(client, make_pdf_pages):
    """P11 — rolls stay text even if they look numeric, so Excel never
    silently reformats/strips a leading zero on later manual edit."""
    pages = ["MBAN301 - Business Mathematics\nRoll No: 00123  Regular\n"]
    with make_pdf_pages(pages):
        files = [("files", ("text_roll.pdf", b"%PDF-1.4 dummy", "application/pdf"))]
        res = await client.post("/api/v1/upload", files=files)
    job_id = res.json()["job_id"]

    res = await client.post("/api/v1/export", json={"job_id": job_id, "filename": "text_roll_out"})
    assert res.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb["Subject-wise Roll Number List"]
    cell = ws.cell(row=3, column=2)
    assert cell.number_format == "@"
    assert cell.data_type == "s"


async def test_repeat_export_reuses_the_same_output_file_row(
    client, test_session_factory, make_pdf_pages
):
    """P1-17 — exporting the same job with the same filename twice must not
    accumulate a duplicate OutputFile row pointing at the same storage key."""
    pages = ["MBAN301 - Business Mathematics\nRoll No: 60001  Regular\n"]
    with make_pdf_pages(pages):
        files = [("files", ("repeat.pdf", b"%PDF-1.4 dummy", "application/pdf"))]
        res = await client.post("/api/v1/upload", files=files)
    job_id = res.json()["job_id"]

    for _ in range(2):
        res = await client.post(
            "/api/v1/export", json={"job_id": job_id, "filename": "repeat_out"}
        )
        assert res.status_code == 200

    async with test_session_factory() as session:
        result = await session.execute(
            select(OutputFile).where(
                OutputFile.job_id == job_id, OutputFile.filename == "output_repeat_out.xlsx"
            )
        )
        rows = result.scalars().all()

    assert len(rows) == 1
